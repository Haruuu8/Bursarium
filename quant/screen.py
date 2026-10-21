"""
AUTO SCREENER - filter otomatis per bulan + RATING.
- Exit: TP High-1% | SL -3% (sudah termasuk fee)
- Tanggal = entry (T+1).
- Minggu dikelompokkan per kalender (Senin-Jumat).
- Sinyal akhir bulan lalu yang entry-nya masuk bulan ini tetap ke-capture.
- Saham suspended (gap > 7 hari) otomatis di-skip.
- STEP 3: auto-regenerate JSON untuk Watchlist.
- STEP 4: auto-push ke GitHub.
"""

import os
import subprocess
import sys
from datetime import datetime
import numpy as np
import pandas as pd

from .data_loader import load_panel_from_gabungan
from .data_cleaner import clean_panel
from .feature_engineering import build_features
from .event_detector import detect_events


MIN_GAIN_FOR_TP = 0.01
EXIT_BELOW_HIGH_PCT = 0.01
SL_PCT = -0.03
MAX_GAP_DAYS = 7
OHLCV_DIR = "data_ohlcv_idx"
PARQUET_FILE = "ohlcv_idx_2026_all.parquet"


# ============================================================
# UTIL
# ============================================================
def get_bulan_berjalan():
    today = pd.Timestamp(datetime.now().date())
    start = today.replace(day=1)
    if today.month == 12:
        end = today.replace(year=today.year + 1, month=1, day=1) - pd.Timedelta(days=1)
    else:
        end = today.replace(month=today.month + 1, day=1) - pd.Timedelta(days=1)
    return start, end, today.strftime("%B %Y"), today.strftime("%Y-%m")


def _week_of_month(dt):
    """Minggu kalender: minggu 1 = minggu yang mengandung tanggal 1 (mulai Senin)."""
    dt = pd.Timestamp(dt)
    first = dt.replace(day=1)
    monday_week1 = first - pd.Timedelta(days=first.weekday())
    delta = (dt - monday_week1).days
    return (delta // 7) + 1


def _periode_minggu(tanggal_list, minggu_ke):
    sample = pd.Timestamp(tanggal_list[0])
    first = sample.replace(day=1)
    monday_week1 = first - pd.Timedelta(days=first.weekday())
    senin = monday_week1 + pd.Timedelta(days=(minggu_ke - 1) * 7)
    jumat = senin + pd.Timedelta(days=4)
    return senin, jumat


def _subprocess_env():
    """Environment untuk subprocess - paksa UTF-8 biar emoji tidak crash di Windows."""
    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONUTF8"] = "1"
    env["PYTHONLEGACYWINDOWSSTDIO"] = "0"
    return env


def jalankan_update():
    print(f"\n{'='*100}")
    print("STEP 1: Update data dari IDX")
    print(f"{'='*100}\n")

    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    update_script = os.path.join(root, "update.py")
    gabung_script = os.path.join(root, "gabung.py")
    parquet_path = os.path.join(root, PARQUET_FILE)

    if not os.path.exists(update_script):
        print(f"[!] File {update_script} tidak ditemukan.")
        return False

    env = _subprocess_env()

    # ---------- Step 1a: update.py ----------
    print("Menjalankan update.py...")
    try:
        result = subprocess.run(
            [sys.executable, update_script],
            cwd=root, capture_output=True, text=True,
            encoding="utf-8", errors="replace",
            env=env, timeout=1800,
        )
        if result.stdout:
            print(result.stdout[-1500:])
        if result.returncode != 0:
            print(f"[!] update.py GAGAL (exit {result.returncode}):")
            if result.stderr:
                print(result.stderr[-800:])
            return False
        print("[OK] update.py selesai")
    except subprocess.TimeoutExpired:
        print("[!] update.py timeout (>30 menit). Proses dibatalkan.")
        return False
    except Exception as e:
        print(f"[!] Gagal jalankan update.py: {e}")
        return False

    # ---------- Step 1b: gabung.py ----------
    if not os.path.exists(gabung_script):
        print(f"[!] File {gabung_script} tidak ditemukan.")
        return False

    print("\nMenjalankan gabung.py...")
    try:
        result = subprocess.run(
            [sys.executable, gabung_script],
            cwd=root, capture_output=True, text=True,
            encoding="utf-8", errors="replace",
            env=env, timeout=900,
        )
        if result.stdout:
            print(result.stdout[-800:])
        if result.returncode != 0:
            print(f"[!] gabung.py GAGAL (exit {result.returncode}):")
            if result.stderr:
                print(result.stderr[-800:])
            return False
        print("[OK] gabung.py selesai")
    except subprocess.TimeoutExpired:
        print("[!] gabung.py timeout (>15 menit). Proses dibatalkan.")
        return False
    except Exception as e:
        print(f"[!] Gagal jalankan gabung.py: {e}")
        return False

    # ---------- Step 1c: verifikasi parquet ----------
    if not os.path.exists(parquet_path):
        print(f"[!] Parquet {PARQUET_FILE} tidak ditemukan setelah gabung.py.")
        return False

    try:
        df = pd.read_parquet(parquet_path)
        last_date = pd.to_datetime(df["Date"]).max()
        print(f"\n[OK] Parquet terakhir: {last_date.date()} ({len(df):,} baris)")
    except Exception as e:
        print(f"[!] Gagal baca parquet: {e}")
        return False

    print("\n[OK] Update + gabung selesai.\n")
    return True


def jalankan_scan_json():
    """STEP 3: regenerate JSON untuk Watchlist."""
    print(f"\n{'#'*100}")
    print(f"#  STEP 3: Regenerate JSON untuk Watchlist")
    print(f"{'#'*100}\n")

    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    scan_script = os.path.join(root, "scan_all.py")

    if not os.path.exists(scan_script):
        print(f"\n[!] ERROR FATAL: File 'scan_all.py' TIDAK DITEMUKAN di folder utama.")
        print(f"    Script ini wajib ada untuk meng-update data_setups.json & data_scans/")
        print(f"    Tanpa file ini, Tab Buy Support, Breakout, Bottom Fishing, & Pre-Cross akan STUCK.\n")
        return False

    env = _subprocess_env()

    print("Menjalankan scan_all.py (5-10 menit)...")
    try:
        result = subprocess.run(
            [sys.executable, scan_script],
            cwd=root, capture_output=True, text=True,
            encoding="utf-8", errors="replace",
            env=env, timeout=1800,
        )
        if result.stdout:
            print(result.stdout[-2500:])
        if result.returncode != 0:
            print(f"[!] scan_all.py exit code {result.returncode}")
            if result.stderr:
                print(result.stderr[-800:])
            return False
        print("\n[OK] Scan JSON selesai.")
        return True
    except subprocess.TimeoutExpired:
        print("[!] scan_all.py timeout (>30 menit). Proses dibatalkan.")
        return False
    except Exception as e:
        print(f"[!] Gagal jalankan scan_all.py: {e}")
        return False


def jalankan_git_push():
    """STEP 4: Auto push ke GitHub."""
    print(f"\n{'#'*100}")
    print(f"#  STEP 4: Auto Push ke GitHub")
    print(f"{'#'*100}\n")

    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    env = _subprocess_env()

    # 1. Git Add
    print("Menjalankan: git add .")
    try:
        subprocess.run(["git", "add", "."], cwd=root, check=True, env=env)
    except Exception as e:
        print(f"[!] Git add gagal: {e}")
        return False

    # 2. Git Commit
    commit_msg = f"Auto-update data: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    print(f"Menjalankan: git commit -m \"{commit_msg}\"")
    try:
        result_commit = subprocess.run(
            ["git", "commit", "-m", commit_msg],
            cwd=root, capture_output=True, text=True, env=env
        )
        
        if result_commit.returncode != 0:
            if "nothing to commit" in result_commit.stdout or "nothing to commit" in result_commit.stderr:
                print("  [i] Tidak ada perubahan data baru untuk di-commit.")
            else:
                print(f"  [!] Git commit gagal (bukan karena kosong):")
                print(f"      {result_commit.stderr}")
                return False
    except Exception as e:
        print(f"[!] Git commit error: {e}")
        return False

    # 3. Git Push
    print("Menjalankan: git push origin main")
    try:
        result_push = subprocess.run(
            ["git", "push", "origin", "main"],
            cwd=root, capture_output=True, text=True, env=env
        )
        if result_push.returncode != 0:
            print(f"  [!] Git push gagal:")
            print(result_push.stderr)
            return False
        
        print("  [OK] Push ke GitHub berhasil! Streamlit akan otomatis me-reboot.")
        return True
    except Exception as e:
        print(f"[!] Git push error: {e}")
        return False


# ============================================================
# RULES
# ============================================================
def rule_a(d):
    return (
        (d["ret_1d"] > 0.15) &
        (d["ret_5d"] > 0.15) &
        (d["px_vs_sma20"] > 0.20)
    )


def rule_b(d):
    return (
        (d["ret_1d"] > 0.15) &
        (d["ret_5d"] > 0.20) &
        (d["range"] > 0.10) &
        (d["vol_z_20"] > 2)
    )


def next_trading_day(ts):
    nxt = pd.Timestamp(ts) + pd.Timedelta(days=1)
    while nxt.weekday() >= 5:
        nxt += pd.Timedelta(days=1)
    return nxt


def add_forward(df):
    df = df.copy()
    df = df.sort_values(["kode", "date"]).reset_index(drop=True)
    for h in [1, 2, 3]:
        df[f"open_T{h}"] = df.groupby("kode")["open"].shift(-h)
        df[f"high_T{h}"] = df.groupby("kode")["high"].shift(-h)
        df[f"low_T{h}"] = df.groupby("kode")["low"].shift(-h)
        df[f"close_T{h}"] = df.groupby("kode")["close"].shift(-h)
    df["date_T1"] = df.groupby("kode")["date"].shift(-1)
    return df


# ============================================================
# SIMULASI TRADE
# ============================================================
_OHLCV_CACHE = {}


def _load_ohlcv(kode):
    if kode in _OHLCV_CACHE:
        return _OHLCV_CACHE[kode]

    path = os.path.join(OHLCV_DIR, f"{kode}.csv")
    if not os.path.exists(path):
        _OHLCV_CACHE[kode] = None
        return None

    try:
        df = pd.read_csv(path)
    except Exception:
        _OHLCV_CACHE[kode] = None
        return None

    rename_map = {}
    for col in df.columns:
        cl = str(col).lower().strip()
        if cl in ("open", "openprice", "open_price"): rename_map[col] = "Open"
        elif cl in ("high", "highprice", "high_price"): rename_map[col] = "High"
        elif cl in ("low", "lowprice", "low_price"): rename_map[col] = "Low"
        elif cl in ("close", "closeprice", "close_price"): rename_map[col] = "Close"
        elif cl in ("date", "tanggal"): rename_map[col] = "Date"

    df = df.rename(columns=rename_map)
    df = df.loc[:, ~df.columns.duplicated(keep="first")]

    if any(c not in df.columns for c in ["Date", "Open", "High", "Low", "Close"]):
        _OHLCV_CACHE[kode] = None
        return None

    try:
        df["Date"] = pd.to_datetime(df["Date"])
        df = df.sort_values("Date").reset_index(drop=True)
        for c in ["Open", "High", "Low", "Close"]:
            df[c] = pd.to_numeric(df[c], errors="coerce")
    except Exception:
        _OHLCV_CACHE[kode] = None
        return None

    _OHLCV_CACHE[kode] = df
    return df


def simulate_trade(kode, entry, tanggal_signal, max_gap_days=MAX_GAP_DAYS):
    df = _load_ohlcv(kode)
    if df is None:
        return None, "SUSPENDED"

    try:
        entry = float(entry)
        sig_dt = pd.Timestamp(tanggal_signal)
    except (ValueError, TypeError):
        return None, "SUSPENDED"

    if pd.isna(entry) or entry <= 0:
        return None, "SUSPENDED"

    m = df.index[df["Date"] == sig_dt].tolist()
    if m:
        sig_idx = m[0]
    else:
        v = df[df["Date"] >= sig_dt]
        if len(v) == 0:
            return None, "SUSPENDED"
        sig_idx = v.index[0]

    entry_idx = sig_idx + 1
    if entry_idx >= len(df):
        return None, "PENDING"

    sig_date = df.iloc[sig_idx]["Date"]
    entry_date = df.iloc[entry_idx]["Date"]
    gap_days = (entry_date - sig_date).days
    if gap_days > max_gap_days:
        return None, "SUSPENDED"

    signal_close = df.iloc[sig_idx]["Close"]
    open_t1 = df.iloc[entry_idx]["Open"]
    if pd.isna(open_t1) or open_t1 <= 0:
        return None, "SUSPENDED"
    if not pd.isna(signal_close) and open_t1 < signal_close:
        return None, "SKIP_GAP_DOWN"

    sl_price = entry * (1 + SL_PCT)

    for i in range(entry_idx, len(df)):
        row = df.iloc[i]
        high = row["High"]
        low = row["Low"]
        if pd.isna(high) or pd.isna(low):
            continue

        high_gain = high / entry - 1
        if high_gain >= MIN_GAIN_FOR_TP:
            exit_price = high * (1 - EXIT_BELOW_HIGH_PCT)
            return exit_price / entry - 1, "TP"

        if low <= sl_price:
            return SL_PCT, "SL"

    last_close = float(df.iloc[-1]["Close"])
    if pd.isna(last_close):
        return None, "SUSPENDED"
    return last_close / entry - 1, "CLOSE"


# ============================================================
# RATING
# ============================================================
def hitung_rating_kandidat(row):
    score = 0
    match = str(row.get("match", "") or "")
    n_match = len(match)
    if n_match >= 2: score += 40
    elif n_match == 1: score += 25

    ret1d = row.get("ret_1d")
    if pd.notna(ret1d):
        if ret1d > 0.25: score += 20
        elif ret1d > 0.20: score += 16
        elif ret1d > 0.17: score += 12
        elif ret1d > 0.15: score += 8

    ret5d = row.get("ret_5d")
    if pd.notna(ret5d):
        if ret5d > 0.30: score += 20
        elif ret5d > 0.20: score += 16
        elif ret5d > 0.15: score += 12
        elif ret5d > 0.10: score += 8

    rng = row.get("range")
    volz = row.get("vol_z_20")
    if pd.notna(rng):
        if rng > 0.15: score += 10
        elif rng > 0.10: score += 6
        elif rng > 0.08: score += 3
    if pd.notna(volz):
        if volz > 3: score += 10
        elif volz > 2: score += 6
        elif volz > 1.5: score += 3

    return round(min(score, 100), 1)


def klasifikasi_rating(score):
    if score >= 75: return "ELITE"
    elif score >= 60: return "KUAT"
    elif score >= 45: return "SEDANG"
    elif score >= 30: return "LEMAH"
    return "HINDARI"


# ============================================================
# EXCEL STYLE
# ============================================================
def excel_style(ws, title_text, subtitle_text, header_row=4):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    ws.merge_cells(start_row=1, start_column=1, end_row=1,
                   end_column=ws.max_column)
    tc = ws.cell(row=1, column=1)
    tc.value = title_text
    tc.font = Font(bold=True, size=16, color="1F4E78")
    tc.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 25

    ws.merge_cells(start_row=2, start_column=1, end_row=2,
                   end_column=ws.max_column)
    sc = ws.cell(row=2, column=1)
    sc.value = subtitle_text
    sc.font = Font(italic=True, size=10, color="666666")
    sc.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 15

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78",
                               fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )
    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[header_row].height = 30
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate
    return thin_border


def autofit_columns(ws, skip_rows=3):
    from openpyxl.utils import get_column_letter
    for col_idx, col in enumerate(ws.iter_cols(min_row=skip_rows + 1,
                                                 max_row=ws.max_row), 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 25)


# ============================================================
# SIMPAN EXCEL
# ============================================================
def simpan_excel(kandidat, signals_bulan, pred_date, path_dir,
                 bulan_label, bulan_file, bulan_start):
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        from openpyxl.formatting.rule import CellIsRule
    except ImportError:
        print("  [!] openpyxl tidak terinstall.")
        return

    filepath = os.path.join(path_dir, f"Rekap_Trading_{bulan_file}.xlsx")

    # SHEET 1: KANDIDAT
    if not kandidat.empty:
        cols = ["rating", "class", "kode", "close", "value",
                "ret_1d", "ret_5d", "range", "vol_z_20",
                "px_vs_sma20", "match"]
        cols = [c for c in cols if c in kandidat.columns]
        df_kand = kandidat[cols].copy()

        for c in ["ret_1d", "ret_5d", "range", "px_vs_sma20"]:
            if c in df_kand.columns:
                df_kand[c] = (df_kand[c] * 100).round(2)
        if "vol_z_20" in df_kand.columns:
            df_kand["vol_z_20"] = df_kand["vol_z_20"].round(2)
        if "value" in df_kand.columns:
            df_kand["value"] = (df_kand["value"] / 1e9).round(2)
        if "rating" in df_kand.columns:
            df_kand["rating"] = df_kand["rating"].round(1)

        df_kand = df_kand.rename(columns={
            "rating": "Rating", "class": "Class",
            "kode": "Kode", "close": "Close", "value": "Value (M)",
            "ret_1d": "Ret1d%", "ret_5d": "Ret5d%",
            "range": "Range%", "vol_z_20": "VolZ",
            "px_vs_sma20": "Px>sma20%", "match": "Match",
        })
        df_kand.insert(0, "No", range(1, len(df_kand) + 1))
    else:
        df_kand = pd.DataFrame({"Info": ["Tidak ada kandidat"]})

    # SHEET 2: REVIEW HARIAN (kolom Tanggal = entry T+1)
    if not signals_bulan.empty:
        review_rows = []
        for _, row in signals_bulan.sort_values("date_T1").iterrows():
            if "date_T1" in row.index and not pd.isna(row["date_T1"]):
                tgl_entry = pd.Timestamp(row["date_T1"])
            else:
                tgl_entry = pd.Timestamp(row["date"])

            ret, status = simulate_trade(
                row["kode"], row["open_T1"], row["date"]
            )

            match = ""
            if row.get("sig_a", False): match += "A"
            if row.get("sig_b", False): match += "B"

            entry = row["open_T1"]
            high = row["high_T1"]

            review_rows.append({
                "Tanggal": tgl_entry.date(),
                "Hari": ['Senin', 'Selasa', 'Rabu', 'Kamis', 'Jumat',
                         'Sabtu', 'Minggu'][tgl_entry.weekday()],
                "Kode": row["kode"],
                "Close T": round(row["close"], 0),
                "Entry": round(entry, 0) if not pd.isna(entry) else "",
                "High T+1": round(high, 0) if not pd.isna(high) else "",
                "Match": match,
                "Status": status,
                "P/L %": round(ret * 100, 2) if ret is not None else "",
            })
        df_review = pd.DataFrame(review_rows)
    else:
        df_review = pd.DataFrame({"Info": [f"Belum ada sinyal di {bulan_label}"]})

    # SHEET 3: REVIEW MINGGUAN
    minggu_rows = []
    if not signals_bulan.empty:
        sr = signals_bulan.copy()
        if "date_T1" in sr.columns:
            sr["tanggal"] = pd.to_datetime(sr["date_T1"])
        else:
            sr["tanggal"] = pd.to_datetime(sr["date"])

        sr["minggu_ke"] = sr["tanggal"].apply(_week_of_month)

        for wk in sorted(sr["minggu_ke"].unique()):
            wk_data = sr[sr["minggu_ke"] == wk]
            n_sig = len(wk_data)
            n_exec = n_win = n_skip = n_pending = n_tp = n_sl = n_close = n_susp = 0
            pl_total = 0

            for _, row in wk_data.iterrows():
                ret, status = simulate_trade(
                    row["kode"], row["open_T1"], row["date"]
                )
                if status == "SKIP_GAP_DOWN":
                    n_skip += 1
                elif status == "SUSPENDED":
                    n_susp += 1
                elif status == "PENDING":
                    n_pending += 1
                else:
                    n_exec += 1
                    pl_total += ret
                    if ret > 0: n_win += 1
                    if status == "TP": n_tp += 1
                    elif status == "SL": n_sl += 1
                    elif status == "CLOSE": n_close += 1

            tgl_sample = wk_data["tanggal"].min()
            senin, jumat = _periode_minggu([tgl_sample], wk)
            tgl_min = wk_data["tanggal"].min().date()
            tgl_max = wk_data["tanggal"].max().date()

            minggu_rows.append({
                "Minggu Ke": f"Minggu {wk}",
                "Senin": senin.date(),
                "Jumat": jumat.date(),
                "Rentang Aktual": f"{tgl_min} s/d {tgl_max}",
                "Sinyal": n_sig,
                "Eksekusi": n_exec,
                "Profit": n_win,
                "TP": n_tp,
                "SL": n_sl,
                "Close": n_close,
                "Skip": n_skip,
                "Suspended": n_susp,
                "Pending": n_pending,
                "Win Rate %": round(n_win / n_exec * 100, 1) if n_exec > 0 else 0,
                "Total P/L %": round(pl_total * 100, 2),
                "Avg P/L %": round(pl_total / n_exec * 100, 2) if n_exec > 0 else 0,
            })

    if minggu_rows:
        df_minggu = pd.DataFrame(minggu_rows)
    else:
        df_minggu = pd.DataFrame({
            "Info": [f"Belum ada sinyal mingguan di {bulan_label}"]
        })

    # SHEET 4: RINGKASAN
    if not signals_bulan.empty:
        total_pl = n_exec = n_win = 0
        n_skip = n_pending = n_tp = n_sl = n_close = n_susp = 0

        for _, row in signals_bulan.iterrows():
            ret, status = simulate_trade(
                row["kode"], row["open_T1"], row["date"]
            )
            if status == "SKIP_GAP_DOWN":
                n_skip += 1
            elif status == "SUSPENDED":
                n_susp += 1
            elif status == "PENDING":
                n_pending += 1
            else:
                n_exec += 1
                total_pl += ret
                if ret > 0: n_win += 1
                if status == "TP": n_tp += 1
                elif status == "SL": n_sl += 1
                elif status == "CLOSE": n_close += 1

        win_rate = (n_win / n_exec * 100) if n_exec > 0 else 0
        avg_pl = (total_pl / n_exec * 100) if n_exec > 0 else 0

        ringkasan_rows = [
            ["PERIODE", f"Bulan {bulan_label}"],
            ["Exit Rule", f"TP High-1% | SL {SL_PCT*100:.0f}% (incl. fee)"],
            ["Tanggal", "Entry (T+1)"],
            ["", ""],
            ["RINGKASAN SINYAL", ""],
            ["Total Sinyal", len(signals_bulan)],
            ["Tereksekusi", n_exec],
            ["  TP (high -1%)", n_tp],
            ["  SL", n_sl],
            ["  Close (fallback)", n_close],
            ["Skip (gap down)", n_skip],
            ["Suspended", n_susp],
            ["Pending", n_pending],
            ["", ""],
            ["PERFORMA", ""],
            ["Win Rate", f"{win_rate:.1f}%"],
            ["Total P/L", f"{total_pl*100:+.2f}%"],
            ["Avg P/L per Trade", f"{avg_pl:+.2f}%"],
            ["", ""],
            ["VERDICT", ""],
            ["Status", "[OK] LAYAK" if win_rate >= 60 and avg_pl > 2
             else ("[!] MARGINAL" if win_rate >= 50 else "[X] EVALUASI")],
        ]
    else:
        ringkasan_rows = [
            ["PERIODE", f"Bulan {bulan_label}"],
            ["Info", "Belum ada sinyal"],
        ]

    df_ringkasan = pd.DataFrame(ringkasan_rows, columns=["Metrik", "Nilai"])

    # TULIS EXCEL
    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        df_ringkasan.to_excel(writer, sheet_name="Ringkasan",
                               index=False, startrow=3)
        df_kand.to_excel(writer, sheet_name="Kandidat Besok",
                          index=False, startrow=3)
        df_minggu.to_excel(writer, sheet_name="Review Mingguan",
                            index=False, startrow=3)
        df_review.to_excel(writer, sheet_name="Review Harian",
                            index=False, startrow=3)

    # FORMATTING
    wb = load_workbook(filepath)

    GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    GREEN_FONT = Font(color="006100", bold=True)
    RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    RED_FONT = Font(color="9C0006", bold=True)
    SECTION_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    SECTION_FONT = Font(bold=True, size=11, color="1F4E78")

    ws = wb["Ringkasan"]
    excel_style(ws, f"RINGKASAN TRADING - {bulan_label.upper()}",
                f"Dibuat: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=2):
        c = row[0]
        if c.value and isinstance(c.value, str):
            if c.value.isupper() and c.value.strip():
                c.fill = SECTION_FILL
                c.font = SECTION_FONT
                row[1].fill = SECTION_FILL
    autofit_columns(ws)

    ws = wb["Kandidat Besok"]
    excel_style(ws, "KANDIDAT BESOK", f"Prediksi untuk: {pred_date}")
    autofit_columns(ws)
    if not kandidat.empty:
        for row in ws.iter_rows(min_row=5, max_row=ws.max_row,
                                 min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.alignment = Alignment(horizontal="center")

    ws = wb["Review Mingguan"]
    excel_style(ws, f"REVIEW MINGGUAN - {bulan_label.upper()}",
                f"Minggu kalender Senin-Jumat - Bulan {bulan_label}")
    autofit_columns(ws)
    if minggu_rows:
        for i, cell in enumerate(ws[4], 1):
            if cell.value in ("Total P/L %", "Avg P/L %"):
                cl = get_column_letter(i)
                rng = f"{cl}5:{cl}{ws.max_row}"
                ws.conditional_formatting.add(rng,
                    CellIsRule(operator="greaterThan", formula=["0"],
                               fill=GREEN_FILL, font=GREEN_FONT))
                ws.conditional_formatting.add(rng,
                    CellIsRule(operator="lessThan", formula=["0"],
                               fill=RED_FILL, font=RED_FONT))

    ws = wb["Review Harian"]
    excel_style(ws, f"REVIEW HARIAN - {bulan_label.upper()}",
                f"TP High-1% | SL {SL_PCT*100:.0f}% - "
                f"Tanggal = entry (T+1) - Bulan {bulan_label}")
    autofit_columns(ws)
    if not df_review.empty and "P/L %" in df_review.columns:
        for i, cell in enumerate(ws[4], 1):
            if cell.value == "P/L %":
                cl = get_column_letter(i)
                rng = f"{cl}5:{cl}{ws.max_row}"
                ws.conditional_formatting.add(rng,
                    CellIsRule(operator="greaterThan", formula=["0"],
                               fill=GREEN_FILL, font=GREEN_FONT))
                ws.conditional_formatting.add(rng,
                    CellIsRule(operator="lessThan", formula=["0"],
                               fill=RED_FILL, font=RED_FONT))

    wb.active = wb.sheetnames.index("Kandidat Besok")
    wb.save(filepath)
    print(f"  -> Excel: {filepath}")


# ============================================================
# MAIN (DENGAN ERROR HANDLING BARU)
# ============================================================
def main():
    print("=" * 100)
    print(f"AUTO SCREENER - {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print("=" * 100)

    bulan_start, bulan_end, bulan_label, bulan_file = get_bulan_berjalan()
    print(f"\nBulan berjalan : {bulan_label}")
    print(f"Periode        : {bulan_start.date()} s/d {bulan_end.date()}")

    # ✅ FIX #1: Cek apakah update berhasil. Kalau gagal, STOP.
    if not jalankan_update():
        print("\n" + "=" * 100)
        print("[!] STOP: Data OHLCV gagal diupdate.")
        print("    Screener dibatalkan agar tidak memproses data lama.")
        print("=" * 100)
        return

    print(f"\n{'='*100}")
    print("STEP 2: Proses screening")
    print(f"{'='*100}\n")

    print("Memuat data...")
    try:
        panel = load_panel_from_gabungan()
    except FileNotFoundError:
        print("[!] File data tidak ditemukan.")
        return

    panel = clean_panel(panel)
    panel = build_features(panel)
    panel = detect_events(panel)
    panel = add_forward(panel)
    panel = panel[(panel["open"] > 0) & (panel["close"] > 0)].copy()

    last_date = pd.Timestamp(panel["date"].max())
    pred_date = next_trading_day(last_date)

    print(f"\nData terakhir  : {last_date.date()} "
          f"({['Senin','Selasa','Rabu','Kamis','Jumat','Sabtu','Minggu'][last_date.weekday()]})")
    print(f"Prediksi untuk : {pred_date.date()} "
          f"({['Senin','Selasa','Rabu','Kamis','Jumat','Sabtu','Minggu'][pred_date.weekday()]})")

    # Kandidat besok
    today_data = panel[panel["date"] == last_date].copy()
    m_a = rule_a(today_data).fillna(False).values
    m_b = rule_b(today_data).fillna(False).values
    m_union = m_a | m_b
    kandidat = today_data[m_union].copy()

    print(f"\n{'#'*100}")
    print(f"#  KANDIDAT UNTUK {pred_date.date()}")
    print(f"{'#'*100}")

    if not kandidat.empty:
        kandidat["match"] = ""
        kandidat.loc[m_a[m_union], "match"] += "A"
        kandidat.loc[m_b[m_union], "match"] += "B"
        kandidat["rating"] = kandidat.apply(hitung_rating_kandidat, axis=1)
        kandidat["class"] = kandidat["rating"].apply(klasifikasi_rating)
        kandidat = kandidat.sort_values(
            ["rating", "value"], ascending=[False, False]
        ).reset_index(drop=True)

        print(f"\n  {len(kandidat)} KANDIDAT:\n")
        print(f"  {'No':>3} {'Kode':<7} {'Rating':>7} {'Class':<12} "
              f"{'Close':>8} {'Ret1d%':>8} {'Ret5d%':>8} "
              f"{'Match':>6} {'Value(M)':>10}")
        print("  " + "-" * 100)
        for i, (_, r) in enumerate(kandidat.iterrows(), 1):
            print(f"  {i:>3} {r['kode']:<7} {r['rating']:>6.1f} "
                  f"{r['class']:<12} {r['close']:>8.0f} "
                  f"{r['ret_1d']*100:>+7.1f}% {r['ret_5d']*100:>+7.1f}% "
                  f"{r['match']:>6} {r['value']/1e9:>9.2f}")
    else:
        print(f"\n  Tidak ada kandidat untuk {pred_date.date()}.")

    # Review bulan
    print(f"\n{'#'*100}")
    print(f"#  REVIEW BULAN {bulan_label.upper()} (tanggal entry)")
    print(f"{'#'*100}")

    _awal = bulan_start - pd.Timedelta(days=10)
    sub_bulan = panel[(panel["date"] >= _awal) &
                       (panel["date"] <= bulan_end)].copy()

    m_a_sub = rule_a(sub_bulan).fillna(False).values
    m_b_sub = rule_b(sub_bulan).fillna(False).values
    m_union_sub = m_a_sub | m_b_sub

    signals_bulan = sub_bulan[m_union_sub].copy()
    signals_bulan["sig_a"] = m_a_sub[m_union_sub]
    signals_bulan["sig_b"] = m_b_sub[m_union_sub]

    if "date_T1" in signals_bulan.columns:
        signals_bulan = signals_bulan[
            (signals_bulan["date_T1"] >= bulan_start) &
            (signals_bulan["date_T1"] <= bulan_end)
        ].copy()

    if not signals_bulan.empty:
        total_pl = n_exec = n_win = 0
        n_skip = n_pending = n_tp = n_sl = n_close = n_susp = 0

        for _, row in signals_bulan.iterrows():
            ret, status = simulate_trade(
                row["kode"], row["open_T1"], row["date"]
            )
            if status == "SKIP_GAP_DOWN":
                n_skip += 1
            elif status == "SUSPENDED":
                n_susp += 1
            elif status == "PENDING":
                n_pending += 1
            else:
                n_exec += 1
                total_pl += ret
                if ret > 0: n_win += 1
                if status == "TP": n_tp += 1
                elif status == "SL": n_sl += 1
                elif status == "CLOSE": n_close += 1

        print(f"\n  Periode        : {bulan_label}")
        print(f"  Total sinyal   : {len(signals_bulan)}")
        print(f"  Skip           : {n_skip}")
        print(f"  Suspended      : {n_susp}")
        print(f"  Pending        : {n_pending}")
        print(f"  Tereksekusi    : {n_exec}")
        print(f"  |- TP          : {n_tp}")
        print(f"  |- SL          : {n_sl}")
        print(f"  - Close        : {n_close}")
        if n_exec > 0:
            print(f"\n  Win rate       : {n_win}/{n_exec} ({n_win/n_exec*100:.1f}%)")
            print(f"  Total P/L      : {total_pl*100:+.2f}%")
            print(f"  Avg per trade  : {total_pl/n_exec*100:+.2f}%")
    else:
        print(f"\n  Belum ada sinyal di bulan {bulan_label}.")

    # Simpan
    print(f"\n{'#'*100}")
    print(f"#  MENYIMPAN EXCEL")
    print(f"{'#'*100}\n")

    os.makedirs("quant/output", exist_ok=True)

    if not kandidat.empty:
        csv_path = f"quant/output/kandidat_{pred_date.date()}.csv"
        cols_csv = ["kode", "close", "value", "ret_1d", "ret_5d",
                    "range", "vol_z_20", "px_vs_sma20", "match", "rating"]
        cols_csv = [c for c in cols_csv if c in kandidat.columns]
        kandidat[cols_csv].to_csv(csv_path, index=False)
        print(f"  -> CSV: {csv_path}")

    simpan_excel(kandidat, signals_bulan, pred_date.date(),
                 "quant/output", bulan_label, bulan_file, bulan_start)

    # ============================================================
    # STEP 3: Regenerate JSON untuk Watchlist
    # ============================================================
    scan_ok = jalankan_scan_json()
    
    # ✅ FIX #2: Peringatan keras kalau scan_all.py gagal
    if not scan_ok:
        print("\n" + "!" * 100)
        print("[!] PERINGATAN BESAR: Gagal regenerate data_setups.json!")
        print("    Watchlist Anda (Buy Support, Breakout, Bottom Fishing, Pre-Cross) TIDAK AKAN UPDATE.")
        print("    Pastikan file 'scan_all.py' ada di folder utama dan tidak error.")
        print("!" * 100 + "\n")

    # ============================================================
    # STEP 4: Auto Push ke GitHub
    # ============================================================
    jalankan_git_push()

    print(f"\n{'='*100}")
    print(f"SELESAI - {datetime.now().strftime('%H:%M:%S')}")
    print(f"{'='*100}")


if __name__ == "__main__":
    main()